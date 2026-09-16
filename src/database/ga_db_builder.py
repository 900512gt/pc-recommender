"""
ga_db_builder.py
從 coolpc_patched.xlsx 建立扁平化的 ga_database.json

入選邏輯：
  1. 型號必須在 coolpc_patched.xlsx 有商品（不在就跳過）
  2. PTT ptt_comment.xlsx 總評論數 + 巴哈 bahamut_matched.xlsx 留言數
     兩者合計 >= 30 才寫入資料庫

使用方式：
  python ga_db_builder.py
"""

import json, re, os
import pandas as pd
import openpyxl
from collections import defaultdict

# ── 路徑設定 ─────────────────────────────────────────
COOLPC_DATA   = "input/coolpc_patched.xlsx"
PTT_XLSX      = "input/ptt_comment.xlsx"
BAHA_XLSX     = "input/bahamut_matched.xlsx"
DB_PATH       = "output/ga_database.json"

PRICE_COL         = "含稅價(NT$)"
COMMENT_THRESHOLD = 30

# ── 類別對照 ──────────────────────────────────────────
SHEET_MAP = {
    "CPU":          ("CPU_統計",   "CPU"),
    "GPU":          ("GPU_統計",   "GPU"),
    "RAM":          ("記憶體_統計", "記憶體"),
    "MB":           ("主機板_統計", "主機板"),
    "SSD":          ("SSD_統計",   "SSD"),
    "HDD":          ("HDD_統計",   "HDD"),
    "AIR_COOLER":   ("風冷_統計",  "風冷"),
    "WATER_COOLER": ("水冷_統計",  "水冷"),
    "CASE":         ("機殼_統計",  "機殼"),
    "PSU":          ("電源_統計",  "電源"),
}

BAHA_CAT_MAP = {
    "CPU": "CPU", "GPU": "GPU", "SSD": "SSD", "HDD": "HDD",
    "主機板": "MB", "記憶體": "RAM", "機殼": "CASE",
    "電源": "PSU", "風冷": "AIR_COOLER", "水冷": "WATER_COOLER",
}

MANUAL_KW   = {
    "DEEPCOOL ASSASSIN IV": "ASSASSIN IV",
    "ADATA XPG S70":        "XPG S70",

    # 原價屋的商品名稱用中文品牌名「貓頭鷹」，我們的型號名卻是英文「Noctua」，
    # 直接比對一定落空。extract_fallback_kw() 也救不了：它把 "NH-D15" 從連字號
    # 拆成 "NH" 與 "D15"，兩個都短於 MIN_KW_LEN=4 被濾掉，最後只剩 "Noctua"
    # 這個關鍵字——而原價屋的名稱裡從頭到尾沒有這個字。
    #
    # 「貓頭鷹 NH-D15」是「貓頭鷹 NH-D15S」的子字串，但 build_model_index() 是
    # 依關鍵字長度由長到短指派、且每筆商品只會被認領一次，所以必須連 D15S 一起
    # 列進來：它的關鍵字較長會先認領自己的商品，D15 才不會把 D15S 一起吃掉。
    #
    # 已知取捨：NH-D15 會連帶認領「貓頭鷹 NH-D15 G2」（新款 8 導管，算同系列但
    # 規格不同）。G2 沒有自己的評論資料所以不是獨立型號，不收的話那幾筆商品會
    # 完全沒人認領，兩害相權取其輕。
    "Noctua NH-D15S": "貓頭鷹 NH-D15S",
    "Noctua NH-D15":  "貓頭鷹 NH-D15",
    "Noctua NH-L9a":  "貓頭鷹 NH-L9a",
}
MIN_KW_LEN  = 4
SKIP_TOKENS = {"TB", "GB", "MB", "W"}
DEFAULT_SOCKETS = ["AM4", "AM5", "LGA1700", "LGA1851"]


# ── 品牌/系列推斷（巴哈限定型號用）──────────────────
def infer_brand_series(cat, model):
    m = model.upper()
    if cat == "CPU":
        if model.startswith("AMD"):
            brand = "AMD"
            for s in ["R9","R7","R5","R3"]:
                if s in m: return brand, s
            return brand, "AMD"
        if model.startswith("Intel"):
            brand = "Intel"
            for s in ["Core Ultra 9","Core Ultra 7","Core Ultra 5","i9","i7","i5","i3"]:
                if s.upper() in m: return brand, s
            return brand, "Intel"
    if cat == "GPU":
        if any(x in m for x in ["RTX","GTX","GT"]):
            for s in ["RTX50","RTX40","RTX30","GTX16","GTX10","GT"]:
                if s in m: return "NVIDIA", s
            return "NVIDIA", "NVIDIA"
        if "RX" in m: return "AMD", "RX"
        if "ARC" in m: return "Intel", "Arc"
    if cat == "RAM":
        for s in ["DDR5","DDR4","DDR3"]:
            if s in m: return s, s
    return "", ""


# ── Step 1：統計各型號評論數（PTT + 巴哈）──────────
def load_comment_counts():
    """
    PTT：從 ptt_comment.xlsx 各類別統計 sheet 讀「總評論數」
    巴哈：從 bahamut_matched.xlsx 明細逐行 count
    回傳 {(cat, model): {"ptt": n, "baha": n, "total": n}}
    """
    counts = defaultdict(lambda: {"ptt": 0, "baha": 0, "total": 0})

    # PTT
    print("  讀取 PTT ptt_comment.xlsx ...")
    wb = openpyxl.load_workbook(PTT_XLSX, read_only=True)
    for cat, (sheet_name, _) in SHEET_MAP.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            model = str(row[0]).strip() if row[0] else ""
            total = row[8]  # 總評論數
            if model and total:
                n = int(total)
                counts[(cat, model)]["ptt"]   += n
                counts[(cat, model)]["total"] += n

    # 巴哈
    print("  讀取 巴哈 bahamut_matched.xlsx ...")
    wb2 = openpyxl.load_workbook(BAHA_XLSX, read_only=True)
    ws2 = wb2["評論資料"]
    for row in ws2.iter_rows(min_row=2, values_only=True):
        cat_raw = str(row[0]).strip() if row[0] else ""
        model   = str(row[1]).strip() if row[1] else ""
        cat     = BAHA_CAT_MAP.get(cat_raw)
        if cat and model:
            counts[(cat, model)]["baha"]  += 1
            counts[(cat, model)]["total"] += 1

    return counts


# ── Step 2：讀取 PTT xlsx 型號資訊（品牌、系列）──────
def load_ptt_model_info():
    info = {}
    wb = openpyxl.load_workbook(PTT_XLSX, read_only=True)
    for cat, (sheet_name, _) in SHEET_MAP.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            model = str(row[0]).strip() if row[0] else ""
            if model:
                info[(cat, model)] = {
                    "brand":  str(row[1]).strip() if row[1] else "",
                    "series": str(row[2]).strip() if row[2] else "",
                }
    return info


# ── 搜尋工具 ──────────────────────────────────────────
def extract_fallback_kw(model):
    parts = model.replace("-", " ").split()
    candidates = [
        p for p in parts
        if len(p) >= MIN_KW_LEN
        and p.upper() not in SKIP_TOKENS
        and not any(p.upper().endswith(s) for s in SKIP_TOKENS)
    ]
    return candidates[-1] if candidates else None


def build_model_index(df, models):
    sorted_models = sorted(models,
                           key=lambda m: len(MANUAL_KW.get(m, m)), reverse=True)
    assigned, result = {}, {m: [] for m in models}
    for m in sorted_models:
        kw = MANUAL_KW.get(m, m)
        matched = df[df["名稱"].str.contains(
            re.escape(kw), na=False, case=False, regex=True)]
        for idx in matched.index:
            if idx not in assigned:
                assigned[idx] = m
                result[m].append(idx)
    return {m: df.loc[idxs] if idxs else pd.DataFrame()
            for m, idxs in result.items()}


def find_matched_rows(df, model, index=None):
    if index is not None:
        rows = index.get(model)
        if rows is not None and not rows.empty:
            return rows
        fallback = extract_fallback_kw(model)
        if fallback:
            matched = df[df["名稱"].str.contains(
                re.escape(fallback), na=False, case=False, regex=True)]
            if not matched.empty:
                return matched
        return None
    kw = MANUAL_KW.get(model, model)
    matched = df[df["名稱"].str.contains(
        re.escape(kw), na=False, case=False, regex=True)]
    if matched.empty:
        fallback = extract_fallback_kw(model)
        if fallback:
            matched = df[df["名稱"].str.contains(
                re.escape(fallback), na=False, case=False, regex=True)]
    return matched if not matched.empty else None


# ── 規格解析器 ───────────────────────────────────────
def val(row, col):
    v = row.get(col)
    if v is None or (isinstance(v, float) and v != v):
        return None
    return v

def to_int(v):     return int(v)           if v is not None else None
def to_float_1(v): return round(float(v), 1) if v is not None else None

def infer_supported_mb(form_factor):
    f = str(form_factor).upper()
    if "E-ATX" in f or "EATX" in f:  return ["ITX","M-ATX","ATX","E-ATX"]
    if "ATX" in f and "M-ATX" not in f and "MATX" not in f:
        return ["ITX","M-ATX","ATX","E-ATX"]
    if "M-ATX" in f or "MATX" in f or "MICRO" in f: return ["ITX","M-ATX"]
    if "ITX" in f:  return ["ITX"]
    return ["ITX","M-ATX","ATX","E-ATX"]

def parse_cpu(row):
    socket = val(row, "腳位")
    return {
        "socket":      socket,
        "tdp":         to_int(val(row, "TDP(W)")),
        "memory_type": "DDR5" if socket in ("AM5","LGA1851") else "DDR4" if socket else None,
        "cores":       to_int(val(row, "核心數")),
        "threads":     to_int(val(row, "執行緒")),
        "base_ghz":    to_float_1(val(row, "基礎時脈(GHz)")),
        "boost_ghz":   to_float_1(val(row, "最大時脈(GHz)")),
        "cache_mb":    to_float_1(val(row, "快取(MB)")),
        "igpu":        val(row, "內顯"),
    }

def parse_gpu(row):
    lc = val(row, "卡長(cm)")
    return {
        "length_mm": int(float(lc)*10) if lc is not None else None,
        "vram_gb":   to_int(val(row, "顯存(GB)")),
        "gpu_chip":  val(row, "GPU晶片"),
        "boost_mhz": to_int(val(row, "Boost(MHz)")),
        "cuda_sp":   to_int(val(row, "CUDA/SP數")),
        "power_conn":val(row, "電源接口"),
    }

def parse_ram(row):
    return {
        "type":        val(row, "規格"),
        "speed":       to_int(val(row, "頻率(MHz)")),
        "capacity_gb": to_int(val(row, "總容量(GB)")),
        "single_gb":   to_int(val(row, "單條容量(GB)")),
        "kit":         val(row, "組合"),
        "cl":          val(row, "CL"),
    }

def parse_mb(row):
    return {
        "socket":      val(row, "腳位"),
        "form_factor": val(row, "板型"),
        "memory_type": val(row, "記憶體規格"),
        "m2_slots":    to_int(val(row, "M.2插槽數")),
        "wifi":        val(row, "內建WiFi"),
    }

def parse_ssd(row):
    return {
        "interface":   val(row, "介面"),
        "capacity_gb": to_int(val(row, "容量(GB)")),
        "form_factor": val(row, "規格"),
        "read_mbs":    to_int(val(row, "讀取(MB/s)")),
        "write_mbs":   to_int(val(row, "寫入(MB/s)")),
        "nand":        val(row, "NAND"),
    }

def parse_hdd(row):
    return {
        "rpm":         to_int(val(row, "轉速(RPM)")),
        "capacity_gb": to_int(val(row, "容量(GB)")),
        "cache_mb":    to_int(val(row, "快取(MB)")),
        "size":        val(row, "尺寸"),
    }

def parse_air_cooler(row):
    hc = val(row, "高度(cm)")
    return {
        "supported_sockets": DEFAULT_SOCKETS,
        "max_tdp":    to_int(val(row, "TDP(W)")),
        "height_mm":  int(float(hc)*10) if hc is not None else None,
        "cooler_type":val(row, "類型"),
        "heat_pipes": to_int(val(row, "導管數")),
    }

def parse_water_cooler(row):
    return {
        "supported_sockets": DEFAULT_SOCKETS,
        "radiator_size": to_int(val(row, "冷排(mm)")),
        "max_tdp":       None,
        "fan_count":     to_int(val(row, "風扇數")),
        "lcd":           val(row, "LCD螢幕"),
    }

def parse_case(row):
    gc = val(row, "顯卡限長(cm)")
    cc = val(row, "散熱器限高(cm)")
    form = val(row, "板型") or ""
    psu_ff = ["SFX","ATX"] if "ITX" in str(form).upper() else ["ATX"]
    return {
        "supported_mb":    infer_supported_mb(form),
        "max_gpu_mm":      int(float(gc)*10) if gc is not None else None,
        "max_cooler_mm":   int(float(cc)*10) if cc is not None else None,
        "max_water_mm":    to_int(val(row, "最大水冷(mm)")),
        "psu_form_factor": psu_ff,
        "tempered_glass":  val(row, "玻璃透側"),
        "front_io":        val(row, "前I/O"),
    }

def parse_psu(row):
    return {
        "wattage":     to_int(val(row, "瓦數(W)")),
        "form_factor": val(row, "規格"),
        "rating":      val(row, "認證"),
        "modular":     val(row, "模組化"),
        "length_cm":   to_float_1(val(row, "長度(cm)")),
        "atx_version": val(row, "ATX版本"),
        "pcie5":       val(row, "PCIe5支援"),
        "warranty_yr": to_int(val(row, "保固(年)")),
    }

PARSERS = {
    "CPU": parse_cpu, "GPU": parse_gpu, "RAM": parse_ram,
    "MB":  parse_mb,  "SSD": parse_ssd, "HDD": parse_hdd,
    "AIR_COOLER":   parse_air_cooler,
    "WATER_COOLER": parse_water_cooler,
    "CASE": parse_case, "PSU": parse_psu,
}


# ── 主程式 ────────────────────────────────────────────
def main():
    print("【Step 1】統計評論數（PTT + 巴哈）...")
    comment_counts = load_comment_counts()
    total_combined = sum(1 for v in comment_counts.values()
                         if v["total"] >= COMMENT_THRESHOLD)
    print(f"  合計 >= {COMMENT_THRESHOLD} 則的型號：{total_combined} 個\n")

    print("【Step 2】讀取 PTT 型號資訊（品牌、系列）...")
    ptt_info = load_ptt_model_info()
    print(f"  共 {len(ptt_info)} 個型號\n")

    coolpc = pd.read_excel(COOLPC_DATA, sheet_name=None)

    db = {}
    total_products      = 0
    skipped_threshold   = 0
    skipped_no_coolpc   = 0

    for cat, (_, coolpc_sheet) in SHEET_MAP.items():
        coolpc_df = coolpc.get(coolpc_sheet)
        if coolpc_df is None:
            print(f"[警告] coolpc 找不到 sheet：{coolpc_sheet}，跳過 {cat}")
            continue

        print(f"\n{'='*60}")
        print(f"  {cat}")
        print(f"{'='*60}")

        db[cat] = []
        parser  = PARSERS[cat]

        # 此類別所有型號，依合計評論數排序
        cat_models = {
            model: info
            for (c, model), info in comment_counts.items()
            if c == cat
        }

        qualified = {m: i for m, i in cat_models.items()
                     if i["total"] >= COMMENT_THRESHOLD}
        qualified = dict(sorted(qualified.items(),
                                key=lambda x: -x[1]["total"]))

        low = {m: i for m, i in cat_models.items()
               if i["total"] < COMMENT_THRESHOLD}
        if low:
            print(f"  ⚠️  評論數不足 {COMMENT_THRESHOLD}，略過 {len(low)} 個型號：")
            for m, i in sorted(low.items(), key=lambda x: -x[1]["total"]):
                print(f"       {m}（PTT:{i['ptt']} 巴哈:{i['baha']} 合計:{i['total']}）")
            skipped_threshold += len(low)

        if not qualified:
            print(f"  此類別無符合門檻的型號")
            continue

        # 預建 coolpc 索引
        model_index = build_model_index(coolpc_df, list(qualified.keys()))

        for model, info in qualified.items():
            # 品牌/系列：優先 PTT xlsx，否則從名稱推斷
            pi     = ptt_info.get((cat, model), {})
            brand  = pi.get("brand")  or infer_brand_series(cat, model)[0]
            series = pi.get("series") or infer_brand_series(cat, model)[1]

            matched = find_matched_rows(coolpc_df, model, index=model_index)

            # 記憶體：過濾筆電規格
            if cat == "RAM" and matched is not None and "類型" in matched.columns:
                matched = matched[matched["類型"] != "筆電"]
                if matched.empty:
                    matched = None

            # coolpc 沒有此型號 → 跳過
            if matched is None:
                print(f"  ⚠️  {model}（合計:{info['total']}則）: coolpc 無此商品，跳過")
                skipped_no_coolpc += 1
                continue

            for _, coolpc_row in matched.iterrows():
                specs = parser(coolpc_row)
                item  = {
                    "ptt_model":     model,
                    "name":          str(coolpc_row["名稱"]).strip(),
                    "brand":         brand,
                    "series":        series,
                    "price":         int(coolpc_row[PRICE_COL]),
                    "comment_count": info["total"],
                    "neg_rate":      None,
                    "score":         None,
                }
                item.update(specs)
                db[cat].append(item)
                total_products += 1

            print(f"  ✅ {model}（PTT:{info['ptt']} 巴哈:{info['baha']} 合計:{info['total']}）"
                  f": {len(matched)} 個商品")

    os.makedirs("output", exist_ok=True)
    with open(DB_PATH, "w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=2)

    print(f"\n{'='*60}")
    print(f"完成")
    print(f"  總商品數：{total_products} 筆")
    print(f"  評論數不足略過：{skipped_threshold} 個型號")
    print(f"  coolpc 無商品略過：{skipped_no_coolpc} 個型號")
    for cat, items in db.items():
        models = len(set(i['ptt_model'] for i in items))
        print(f"  {cat:15s}: {len(items):4d} 筆  {models} 個型號")
    print(f"\n已儲存至 {DB_PATH}")
    print(f"\n下一步：")
    print(f"  python merge_and_score.py")
    print(f"  python ga_db_manager.py score output/scores.csv")


if __name__ == "__main__":
    main()