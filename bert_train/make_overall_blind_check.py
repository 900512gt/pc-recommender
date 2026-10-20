"""
make_overall_blind_check.py — 抽查「整則評論情緒 label」能不能直接當總體面向用

總體面向選項 2：沿用相關性判斷階段 GPT 給每則相關評論的整體傾向（正評／負評／中立）。
這支程式從相關評論隨機抽 PTT、巴哈各 50 則，做成盲測表（不顯示 GPT 的答案），
人工填完後再用 --score 比對一致率。

使用方式（在 bert_train/ 底下）：
  python make_overall_blind_check.py            # 產生 總體label盲測.xlsx 與 總體label盲測_答案.csv
  python make_overall_blind_check.py --score    # 人工填完後計算一致率與錯誤清單

答案檔不要先打開，填完盲測表再比對。
"""
import csv
import json
import random
import argparse
from pathlib import Path
from collections import Counter

HERE = Path(__file__).resolve().parent
BLIND = HERE / "總體label盲測.xlsx"
KEY = HERE / "總體label盲測_答案.csv"
CHOICES = ["正評", "負評", "中立"]

GUIDE = [
    ("任務", "判斷這則評論對「指定型號」的整體態度，填在「你的判斷」欄（下拉選單）。"),
    ("正評", "整體而言對這個型號是好的評價：稱讚、滿意、推薦它、說它某方面好。"),
    ("負評", "整體而言對這個型號是不好的評價：批評、不滿、勸人別買、說它有問題或某方面差。"),
    ("中立", "沒有明顯的好壞態度：提問、規格或價格資訊、新聞轉述、好壞參半說不出結論。"),
    ("注意 1", "評價對象必須是指定型號。評論在誇或罵別的型號，對指定型號就是中立（除非也明確評價了它）。"),
    ("注意 2", "PTT 推文的對象通常要看文章標題；稱讚文章或作者（好文、感謝分享）不算評價型號。"),
    ("注意 3", "依自己的判斷填，不確定就填你覺得最接近的；可以在「備註」欄寫原因。"),
]


def make(args):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = [json.loads(l) for l in open(HERE / "aspects_標註_p3.jsonl", encoding="utf-8")]
    rel = [(i, r) for i, r in enumerate(rows) if r.get("orig_relevant") and r.get("label") in CHOICES]
    rng = random.Random(args.seed)
    picked = []
    for src in ("ptt", "baha"):
        pool = [x for x in rel if x[1]["source"] == src]
        picked += rng.sample(pool, args.per_source)
    rng.shuffle(picked)   # 打亂順序，填的人看不出來源分組

    wb = Workbook()
    ws = wb.active
    ws.title = "盲測"
    header = ["編號", "來源", "類別", "型號", "文章標題", "評論內容", "你的判斷", "備註"]
    ws.append(header)
    for n, (i, r) in enumerate(picked, 1):
        ws.append([n, "PTT" if r["source"] == "ptt" else "巴哈", r["category"], r["model"],
                   r.get("title") or "", r.get("content") or "", "", ""])
    widths = [6, 6, 8, 18, 36, 80, 12, 24]
    for col, w in zip("ABCDEFGH", widths):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{len(picked) + 1}")

    guide = wb.create_sheet("說明", 0)
    for k, v in GUIDE:
        guide.append([k, v])
    guide.column_dimensions["A"].width = 10
    guide.column_dimensions["B"].width = 100
    for row in guide.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True)
    wb.save(BLIND)

    with open(KEY, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["編號", "_index", "來源", "GPT標注"])
        for n, (i, r) in enumerate(picked, 1):
            w.writerow([n, i, r["source"], r["label"]])
    print(f"盲測表：{BLIND}（{len(picked)} 則）\n答案檔：{KEY}（填完盲測表之前不要打開）")


def score(args):
    from openpyxl import load_workbook
    ws = load_workbook(BLIND)["盲測"]
    human = {}
    texts = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        human[int(row[0])] = (row[6] or "").strip()
        texts[int(row[0])] = (row[3], row[5])
    key = {int(r["編號"]): r for r in csv.DictReader(open(KEY, encoding="utf-8-sig"))}
    filled = [n for n in key if human.get(n) in CHOICES]
    if len(filled) < len(key):
        print(f"注意：{len(key) - len(filled)} 則還沒填，只計算已填的 {len(filled)} 則")

    def report(name, ns):
        if not ns:
            return
        agree = sum(human[n] == key[n]["GPT標注"] for n in ns)
        print(f"{name:<4} {agree}/{len(ns)} = {agree / len(ns):.0%}")

    report("全部", filled)
    report("PTT", [n for n in filled if key[n]["來源"] == "ptt"])
    report("巴哈", [n for n in filled if key[n]["來源"] == "baha"])
    print("\n混淆表（GPT → 人工）")
    for (g, h), c in sorted(Counter((key[n]["GPT標注"], human[n]) for n in filled).items()):
        print(f"  {g} → {h}：{c}")
    print("\n不一致的評論")
    for n in filled:
        if human[n] != key[n]["GPT標注"]:
            model, text = texts[n]
            print(f"  #{n} GPT {key[n]['GPT標注']} / 人工 {human[n]}｜{model}｜{(text or '')[:60]}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--score", action="store_true", help="人工填完後計算一致率")
    ap.add_argument("--per-source", type=int, default=50)
    ap.add_argument("--seed", type=int, default=20261020)
    args = ap.parse_args()
    score(args) if args.score else make(args)
